from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "materiais/lista_exercicios_docker.docx"
BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F4F6"
GRID = "D9D9D9"

exercicios = [
    ("01", "Transformando uma aplicação em imagem", "Dockerfile e imagem", "Uma API Flask de calculadora está disponível em uma pasta com main.py e requirements.txt. Ela deve responder na porta 5050.", ["Crie um Dockerfile baseado em uma imagem Python enxuta.", "Defina /app como diretório de trabalho, copie os arquivos e instale as dependências.", "Documente a porta da aplicação e defina o comando para iniciá-la.", "Construa uma imagem chamada calculadora:1.0 e execute um teste local."], "Dockerfile e os comandos usados para construir e executar a imagem."),
    ("02", "Publicando a porta correta", "Portas de container e do computador", "A imagem calculadora:1.0 já existe e a aplicação escuta em 5050 dentro do container. Um colega executou docker run calculadora:1.0, mas não consegue acessá-la pelo navegador.", ["Inicie o container com um nome fácil de identificar.", "Publique a aplicação na porta 3000 do computador.", "Teste a rota GET /rotas com navegador ou curl.", "Explique, em uma frase, a diferença entre EXPOSE e publicar uma porta com -p."], "Comando docker run, resultado do teste e explicação curta."),
    ("03", "Uma rota nova sem mexer no container", "Rebuild de imagem", "A API de calculadora precisa ganhar uma rota POST /potencia que recebe a e b em JSON e devolve o resultado de a elevado a b.", ["Adicione a rota seguindo o padrão das operações existentes.", "Teste a rota diretamente antes de criar uma nova imagem.", "Reconstrua a imagem com uma nova tag.", "Suba um novo container e confirme que /rotas lista /potencia."], "Código da rota, comando de build e evidência do teste."),
    ("04", "Trocando o servidor de desenvolvimento", "Processos e Gunicorn", "O Dockerfile atual inicia Flask com python main.py. Prepare uma versão que use Gunicorn e continue atendendo na porta 5050.", ["Confira se gunicorn está nas dependências.", "Ajuste somente o comando de inicialização do Dockerfile.", "Escolha dois workers e faça o bind em 0.0.0.0:5050.", "Reconstrua e confirme pelos logs qual processo iniciou a aplicação."], "Dockerfile atualizado e trecho dos logs que comprove o Gunicorn."),
    ("05", "Separando frontend e backend", "Redes Docker", "Você recebeu duas imagens: calculadora:2.0 para o backend e calc-front:1.0 para o frontend. O frontend deve ficar disponível em http://localhost:8080 e conversar com o backend sem expor o backend para fora.", ["Crie uma rede Docker chamada calc-net.", "Inicie o backend ligado apenas a essa rede.", "Inicie o frontend na mesma rede e publique apenas a porta 80 do frontend como 8080.", "Explique por que o navegador não deve chamar backend:5050 diretamente."], "Comandos usados, saída de docker network inspect e explicação."),
    ("06", "Entendendo o proxy reverso", "Nginx e caminhos relativos", "O frontend faz fetch para /api/rotas. O Nginx precisa servir os arquivos estáticos e encaminhar tudo que começa com /api/ para o serviço backend.", ["Configure a raiz dos arquivos estáticos em /usr/share/nginx/html.", "Crie o bloco location /api/.", "Faça /api/soma chegar ao backend como /soma.", "Inclua os cabeçalhos Host e X-Forwarded-For e descreva a função de um deles."], "Arquivo nginx.conf e teste de GET /api/rotas pelo frontend."),
    ("07", "Compondo os dois serviços", "Docker Compose", "Converta a solução do exercício anterior em um docker-compose.yml com os serviços frontend e backend.", ["Use os diretórios frontend e backend como contextos de build.", "Publique apenas 8080:80 para o frontend.", "Faça o frontend depender do backend.", "Inicie tudo com um único comando e liste os serviços em execução."], "Arquivo docker-compose.yml e saída de docker compose ps."),
    ("08", "Diagnóstico de erro 502", "Logs e portas internas", "O Compose inicia sem erros aparentes, mas GET /api/rotas devolve HTTP 502. Os logs indicam que o Nginx tentou conectar a backend:5000, enquanto o backend informa que está ouvindo em 5050.", ["Identifique a causa em uma frase.", "Corrija a configuração mínima necessária.", "Recrie apenas o que precisar para aplicar a alteração.", "Teste GET /api/rotas e POST /api/soma com o frontend."], "Antes e depois da linha corrigida, mais os dois resultados de teste."),
    ("09", "Validando entradas e respostas", "HTTP JSON e tratamento de erros", "A API deve rejeitar corpos ausentes, campos a ou b faltando, valores não numéricos e divisão por zero. Respostas de erro devem usar HTTP 400 e JSON com a chave erro.", ["Monte uma tabela de testes com pelo menos cinco cenários: dois válidos e três inválidos.", "Envie as requisições usando curl, Postman ou Insomnia.", "Verifique status HTTP e corpo de cada resposta.", "Ajuste o backend caso algum cenário não siga a regra."], "Tabela preenchida com requisição, status esperado, status obtido e resposta."),
    ("10", "Desafio de integração e entrega", "Compose completo e investigação", "Monte uma versão final da calculadora com frontend Nginx, backend Flask com Gunicorn e Docker Compose. Em seguida, introduza de propósito um erro de comunicação e documente como você o encontrou e corrigiu.", ["Garanta que a página carregue os botões a partir de /api/rotas.", "Execute soma, subtração, multiplicação e divisão; inclua uma divisão por zero.", "Provoque um erro simples no proxy, registre o sintoma e consulte os logs.", "Reverta a falha, valide a aplicação e finalize com docker compose down."], "Pasta do projeto, compose, Dockerfiles, nginx.conf, roteiro de diagnóstico e capturas ou saídas dos testes."),
]

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),GRID); b.append(e)
    tblPr.append(b)

def set_cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(11)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(25); styles['Title'].font.color.rgb = RGBColor(0,0,0)
for name, size in [('Heading 1',17),('Heading 2',13)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Lista de exercícios de Docker')
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=sub.add_run('Aplicações Python, redes, Nginx e Docker Compose'); r.italic=True; r.font.size=Pt(13)
doc.add_paragraph('Esta sequência foi pensada para praticar, de forma gradual, as mesmas habilidades trabalhadas nos casos do repositório: transformar uma API em imagem, publicar portas, operar containers em rede, usar proxy reverso e diagnosticar falhas de integração entre frontend e backend.')
doc.add_paragraph('Orientação de uso: os exercícios são independentes, mas ficam mais proveitosos quando realizados na ordem. Registre os comandos, os arquivos alterados e o resultado de cada teste. Antes de avançar, confirme que a etapa funciona.')

t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
widths=[Inches(.8), Inches(3.55), Inches(2.4)]
for i, text in enumerate(['Exercício','Tema','Competência principal']):
    c=t.rows[0].cells[i]; c.width=widths[i]; shade(c, BLUE); set_cell_text(c,text,True,'FFFFFF',10)
for n, name, topic, _, _, _ in exercicios:
    cells=t.add_row().cells
    for i, val in enumerate([n,name,topic]): cells[i].width=widths[i]; set_cell_text(cells[i],val, size=9)
doc.add_paragraph('')

for idx,(num,titulo,tema,contexto,tarefas,entrega) in enumerate(exercicios):
    doc.add_page_break()
    h=doc.add_paragraph(style='Heading 1'); h.add_run(f'Exercício {num}  {titulo}')
    p=doc.add_paragraph(); rr=p.add_run('Foco: '); rr.bold=True; p.add_run(tema)
    p=doc.add_paragraph(); rr=p.add_run('Cenário. '); rr.bold=True; p.add_run(contexto)
    h2=doc.add_paragraph(style='Heading 2'); h2.add_run('O que fazer')
    for item in tarefas:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(5); p.add_run(item)
    h2=doc.add_paragraph(style='Heading 2'); h2.add_run('Entrega esperada')
    p=doc.add_paragraph(); p.add_run(entrega)
    h2=doc.add_paragraph(style='Heading 2'); h2.add_run('Registro do aluno')
    table=doc.add_table(rows=4, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False; borders(table)
    labels=['Comandos executados','Arquivos alterados','Resultado dos testes','Dúvida ou descoberta']
    for row,label in zip(table.rows,labels):
        row.cells[0].width=Inches(1.45); row.cells[1].width=Inches(5.3); shade(row.cells[0],LIGHT_BLUE); set_cell_text(row.cells[0],label,True, size=9); set_cell_text(row.cells[1],'',size=10)
        row.height = Inches(.43 if label != 'Comandos executados' else .62)

for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Estudo Docker e aplicações  |  Lista de prática').font.size=Pt(8)

doc.save(OUT)
print(OUT)
